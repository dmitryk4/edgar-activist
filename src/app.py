from __future__ import annotations

import json
import re
import subprocess
import sys
import threading
import uuid
from pathlib import Path

import openpyxl
from flask import Flask, Response, jsonify, render_template, request, send_from_directory, stream_with_context

SRC_DIR = Path(__file__).parent
OUTPUT_DIR = SRC_DIR.parent / "data" / "output"

app = Flask(__name__)

# job_id -> {lines, done, file, error, event}
_jobs: dict[str, dict] = {}
_lock = threading.Lock()


def _run_scraper(job_id: str, ticker: str) -> None:
    proc = subprocess.Popen(
        [sys.executable, "main.py", ticker],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        cwd=str(SRC_DIR),
    )

    with _lock:
        job = _jobs[job_id]

    output_file: str | None = None
    for raw_line in proc.stdout:
        line = raw_line.rstrip()
        if line.startswith("Saved:"):
            output_file = Path(line.split("Saved:", 1)[-1].strip()).name
        with _lock:
            job["lines"].append(line)
        job["event"].set()

    proc.wait()

    with _lock:
        job["done"] = True
        job["file"] = output_file
        if proc.returncode != 0 and not output_file:
            job["error"] = f"Process exited with code {proc.returncode}"
    job["event"].set()


VIC_UNIVERSE = [
    {
        "ticker": "ONL",
        "name": "Orion Properties",
        "vic_url": "https://valueinvestorsclub.com/idea/Orion_Properties/0013978503",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001873923&type=&dateb=&owner=include&count=40",
    },
    {
        "ticker": "STHO",
        "name": "Star Holdings",
        "vic_url": "https://valueinvestorsclub.com/idea/Star_Holdings/8294145308",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001953366&type=&dateb=&owner=include&count=40",
    },
    {
        "ticker": "VRA",
        "name": "Vera Bradley",
        "vic_url": "https://valueinvestorsclub.com/idea/Vera_Bradley/3912325634",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001495320&type=&dateb=&owner=include&count=40",
    },
    {
        "ticker": "THRY",
        "name": "Thryv",
        "vic_url": "https://valueinvestorsclub.com/idea/THRYV_HOLDINGS_INC/8764916226",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001556739&type=&dateb=&owner=include&count=40",
    },
    {
        "ticker": "MRSN",
        "name": "Mersana Therapeutics",
        "vic_url": "https://valueinvestorsclub.com/idea/MERSANA_THRPEUTIC_INC/0658619873",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001442836&type=&dateb=&owner=include&count=40",
    },
    {
        "ticker": "SEG",
        "name": "Seaport Entertainment Group",
        "vic_url": "https://valueinvestorsclub.com/idea/SEAPORT_ENTERTAINMENT_GR_INC/6129591525",
        "edgar_url": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0002009684&type=&dateb=&owner=include&count=40",
    },
]


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/vic")
def vic():
    return render_template("vic.html", universe=VIC_UNIVERSE)


@app.route("/sumzero")
def sumzero():
    return render_template("coming_soon.html",
        active="sumzero",
        title="SumZero",
        description="SumZero idea universe and institutional holder mapping — coming soon.",
    )


@app.route("/seekingalpha")
def seekingalpha():
    return render_template("coming_soon.html",
        active="seekingalpha",
        title="Seeking Alpha",
        description="Seeking Alpha Marketplace author coverage and micro-cap ideas — coming soon.",
    )


@app.route("/microcapclub")
def microcapclub():
    return render_template("coming_soon.html",
        active="microcapclub",
        title="Microcap Club",
        description="Microcap Club network and Ian Cassel idea universe — coming soon.",
    )


@app.route("/scrape", methods=["POST"])
def start_scrape():
    ticker = request.form.get("ticker", "").strip().upper()
    if not re.match(r"^[A-Z]{1,7}$", ticker):
        return {"error": "Invalid ticker — use 1–7 letters (e.g. ONL)"}, 400

    job_id = str(uuid.uuid4())
    with _lock:
        _jobs[job_id] = {
            "ticker": ticker,
            "lines": [],
            "done": False,
            "file": None,
            "error": None,
            "event": threading.Event(),
        }

    threading.Thread(target=_run_scraper, args=(job_id, ticker), daemon=True).start()
    return {"job_id": job_id}


@app.route("/stream/<job_id>")
def stream(job_id: str):
    with _lock:
        job = _jobs.get(job_id)
    if not job:
        return {"error": "Job not found"}, 404

    def generate():
        cursor = 0
        while True:
            job["event"].wait(timeout=20)
            job["event"].clear()

            with _lock:
                new_lines = job["lines"][cursor:]
                is_done = job["done"]
                out_file = job["file"]
                err = job["error"]

            cursor += len(new_lines)

            for line in new_lines:
                yield f"data: {json.dumps({'type': 'log', 'message': line})}\n\n"

            if is_done:
                if out_file:
                    yield f"data: {json.dumps({'type': 'done', 'file': out_file})}\n\n"
                else:
                    yield f"data: {json.dumps({'type': 'error', 'message': err or 'Scrape failed'})}\n\n"
                break

            yield f"data: {json.dumps({'type': 'ping'})}\n\n"

    return Response(
        stream_with_context(generate()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/download/<path:filename>")
def download(filename: str):
    safe = Path(filename).name
    return send_from_directory(str(OUTPUT_DIR), safe, as_attachment=True)


@app.route("/results/<job_id>")
def results(job_id: str):
    with _lock:
        job = _jobs.get(job_id)
    if not job or not job.get("file"):
        return jsonify({"error": "Results not ready"}), 404

    ticker   = job["ticker"]
    xl_path  = OUTPUT_DIR / job["file"]

    try:
        wb = openpyxl.load_workbook(xl_path, read_only=True, data_only=True)
        if ticker not in wb.sheetnames:
            return jsonify({"error": f"No sheet for {ticker}"}), 404

        ws      = wb[ticker]
        rows    = list(ws.iter_rows(values_only=True))
        headers = [str(h) if h is not None else "" for h in rows[0]]
        holders = []
        for row in rows[1:]:
            if not any(v for v in row):
                continue
            holders.append(dict(zip(headers, [str(v) if v is not None else "" for v in row])))

        return jsonify({"ticker": ticker, "holders": holders})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 5000))
    print(f"Starting EDGAR scraper UI at http://localhost:{port}")
    app.run(debug=False, host="0.0.0.0", port=port, threaded=True)
